from io import BytesIO

from docx import Document

from opspilot.knowledge.chunking import chunk_blocks
from opspilot.knowledge.parsers.base import BlockKind, ParsedBlock
from opspilot.knowledge.parsers.docx import DocxParser


def test_heading_path_and_table_integrity_are_preserved() -> None:
    blocks = [
        ParsedBlock(BlockKind.HEADING, "3 退款", level=1),
        ParsedBlock(BlockKind.HEADING, "3.2 七天无理由", level=2),
        ParsedBlock(BlockKind.PARAGRAPH, "用户可在签收后七天内申请退款。"),
        ParsedBlock(BlockKind.TABLE, "条件 | 处理\n未拆封 | 全额退款"),
    ]

    chunks = chunk_blocks(blocks, max_tokens=30, overlap_tokens=5)
    target = next(chunk for chunk in chunks if "七天无理由" in chunk.content)

    assert target.section_path == ["3 退款", "3.2 七天无理由"]
    assert target.token_count <= 30
    assert any("条件 | 处理\n未拆封 | 全额退款" in chunk.content for chunk in chunks)


def test_chunks_have_bounded_overlap() -> None:
    blocks = [ParsedBlock(BlockKind.PARAGRAPH, " ".join(f"word{i}" for i in range(30)))]

    chunks = chunk_blocks(blocks, max_tokens=12, overlap_tokens=3)

    assert len(chunks) >= 3
    for previous, current in zip(chunks, chunks[1:], strict=False):
        assert previous.content.split()[-3:] == current.content.split()[:3]


def test_chinese_text_is_counted_and_every_chunk_respects_budget() -> None:
    blocks = [ParsedBlock(BlockKind.PARAGRAPH, "退款政策适用于已签收商品" * 30)]

    chunks = chunk_blocks(blocks, max_tokens=40, overlap_tokens=5)

    assert len(chunks) > 1
    assert all(chunk.token_count <= 40 for chunk in chunks)


def test_small_paragraphs_in_same_section_are_merged() -> None:
    blocks = [
        ParsedBlock(BlockKind.HEADING, "退款", level=1),
        ParsedBlock(BlockKind.PARAGRAPH, "第一段说明。"),
        ParsedBlock(BlockKind.PARAGRAPH, "第二段说明。"),
    ]

    chunks = chunk_blocks(blocks, max_tokens=40, overlap_tokens=5)

    assert len(chunks) == 1
    assert "第一段说明" in chunks[0].content
    assert "第二段说明" in chunks[0].content


def test_oversized_table_is_split_without_exceeding_budget() -> None:
    table = "\n".join(f"第{index}行 | " + "内容" * 20 for index in range(10))

    chunks = chunk_blocks([ParsedBlock(BlockKind.TABLE, table)], max_tokens=30, overlap_tokens=5)

    assert len(chunks) > 1
    assert all(chunk.token_count <= 30 for chunk in chunks)


def test_docx_parser_preserves_interleaved_body_order() -> None:
    document = Document()
    document.add_heading("Refund", level=1)
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Condition"
    table.cell(0, 1).text = "Result"
    document.add_paragraph("After table")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)

    blocks = DocxParser().parse(stream)

    assert [block.kind for block in blocks] == [
        BlockKind.HEADING,
        BlockKind.PARAGRAPH,
        BlockKind.TABLE,
        BlockKind.PARAGRAPH,
    ]
